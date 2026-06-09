from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "CodeStory_AI助手模块需求文档_V1.0.docx"

INK = "162033"
BLUE = "2457C5"
DARK_BLUE = "173B72"
LIGHT_BLUE = "EAF1FF"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
GREEN = "16794A"
LIGHT_GREEN = "EAF7F0"
GOLD = "8A5A00"
LIGHT_GOLD = "FFF5D6"
RED = "A52A2A"
LIGHT_RED = "FDECEC"
WHITE = "FFFFFF"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=11, bold=None, color=INK, italic=None, font="Microsoft YaHei"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.15):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MID_GRAY)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.82)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color, before, after in (
        ("Title", 25, BLACK, 0, 5),
        ("Subtitle", 12, MID_GRAY, 0, 14),
        ("Heading 1", 16, BLUE, 16, 7),
        ("Heading 2", 13, DARK_BLUE, 12, 5),
        ("Heading 3", 11.5, DARK_BLUE, 9, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.line_spacing = 1.15

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, after=0)
    run = p.add_run("CODESTORY  /  PRODUCT REQUIREMENTS")
    set_run_font(run, size=8.5, bold=True, color=MID_GRAY)

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_field(p)


def add_title_block(doc):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=10, after=2)
    run = p.add_run("CODESTORY")
    set_run_font(run, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph(style="Title")
    run = p.add_run("AI 助手模块需求文档")
    set_run_font(run, size=25, bold=True, color=BLACK)

    p = doc.add_paragraph(style="Subtitle")
    run = p.add_run("面向编程学习的 AI 导师、混合判题与个性化反馈")
    set_run_font(run, size=12, color=MID_GRAY)

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    rows = [
        ("文档版本", "V1.0"),
        ("文档状态", "方向确认稿"),
        ("产品模块", "小节学习 / AI 助手 / 练习判题"),
        ("编制日期", "2026-06-09"),
        ("适用对象", "产品设计、前端、后端、AI 工程、测试"),
    ]
    for idx, (label, value) in enumerate(rows):
        set_cell_shading(table.cell(idx, 0), LIGHT_BLUE)
        p1 = table.cell(idx, 0).paragraphs[0]
        p2 = table.cell(idx, 1).paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(label)
        set_run_font(r1, size=9.5, bold=True, color=DARK_BLUE)
        r2 = p2.add_run(value)
        set_run_font(r2, size=9.5, color=INK)
    set_table_geometry(table, [1900, 7460])

    add_callout(
        doc,
        "核心决策",
        "选择题继续使用数据库标准答案确定性判分；编程题采用“隔离测试执行 + AI 代码质量评审”的混合评分。测试结果决定功能正确性，AI 负责可读性、实现合理性、效率与教学反馈。",
        LIGHT_GREEN,
        GREEN,
    )


def add_callout(doc, label, text, fill=LIGHT_BLUE, color=DARK_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=1, after=1, line=1.15)
    run = p.add_run(f"{label}  ")
    set_run_font(run, size=10.5, bold=True, color=color)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=INK)
    set_table_geometry(table, [9360])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    set_paragraph_spacing(p)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        set_cell_shading(hdr.cells[i], header_fill)
        p = hdr.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.05)
        run = p.add_run(header)
        set_run_font(run, size=font_size, bold=True, color=DARK_BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, after=0, line=1.08)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    size = {1: 16, 2: 13, 3: 11.5}[level]
    color = BLUE if level == 1 else DARK_BLUE
    set_run_font(r, size=size, bold=True, color=color)
    return p


def add_page_break(doc):
    doc.add_page_break()


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_document(doc)
    add_title_block(doc)

    add_heading(doc, "1. 文档说明", 1)
    add_heading(doc, "1.1 编写目的", 2)
    add_body(
        doc,
        "本文档用于定义 CodeStory AI 助手模块的产品目标、功能范围、核心流程、混合判题策略、AI 行为约束、数据要求与验收标准。本文档只描述“要实现什么”和“达到什么效果”，不限定具体代码结构与模型供应商。",
    )

    add_heading(doc, "1.2 背景与现状", 2)
    add_body(
        doc,
        "当前平台已经具备课程、章节、小节、选择题、编程题、分级提示、用户答案、学习进度和 AI 会话状态等基础数据。现有选择题通过数据库答案比对完成判分；编程题则在规范化代码后与标准答案进行字符串比较。",
    )
    add_body(
        doc,
        "字符串比较无法识别多种等价实现。例如学生使用不同循环结构、变量名或算法实现时，即使结果正确，也可能被判错。因此 AI 助手模块需要升级编程题评价方式，同时保持结果可靠、评分可解释、学习过程可引导。",
    )

    add_heading(doc, "1.3 术语定义", 2)
    add_table(
        doc,
        ["术语", "定义"],
        [
            ("确定性判分", "由数据库标准答案或测试用例产生可重复的正确/错误结果。"),
            ("功能正确性", "学生代码能否编译或解释执行，并通过公开及隐藏测试用例。"),
            ("质量评分", "对可读性、结构、复杂度、规范性等非功能维度进行评价。"),
            ("AI 导师", "基于当前课程、小节、练习和用户进度提供讲解、追问、提示与反馈的助手。"),
            ("Rubric", "固定的评分维度、分值上限、扣分条件和输出约束。"),
            ("判题沙箱", "限制资源、网络和系统权限的隔离代码执行环境。"),
        ],
        [1900, 7460],
    )

    add_page_break(doc)
    add_heading(doc, "2. 产品定位与目标", 1)
    add_heading(doc, "2.1 产品定位", 2)
    add_callout(
        doc,
        "产品定位",
        "AI 助手不是通用聊天机器人，而是嵌入小节学习流程的编程导师。它应帮助学生理解、尝试、纠错和复盘，而不是直接代替学生完成练习。",
    )

    add_heading(doc, "2.2 业务目标", 2)
    for text in (
        "提升编程题判分的合理性，允许存在多个正确实现。",
        "为错误代码提供具体、可执行、符合当前知识水平的反馈。",
        "将“直接给答案”改造成逐级提示和引导式学习。",
        "基于提交、提示和学习进度积累可解释的掌握度数据。",
        "为后续个性化练习、薄弱知识点分析和任务式学习奠定基础。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "2.3 成功指标", 2)
    add_table(
        doc,
        ["指标", "目标口径", "V1 建议目标"],
        [
            ("判题一致性", "同一份代码与同一测试集重复提交的功能结果一致", "100%"),
            ("AI 输出可用率", "返回结构符合约定且可被系统解析", ">= 99%"),
            ("反馈相关性", "反馈能对应实际失败测试或代码问题", "人工抽检 >= 85%"),
            ("平均反馈时延", "提交到展示完整评价的时间", "<= 8 秒"),
            ("直接答案抑制", "未达到最高提示等级前不输出完整标准答案", ">= 95%"),
            ("评分可追溯", "每次评分可关联测试结果、Rubric 和模型版本", "100%"),
        ],
        [1700, 4840, 2820],
    )

    add_heading(doc, "2.4 非目标", 2)
    for text in (
        "V1 不建设支持任意语言和任意依赖的通用在线 IDE。",
        "V1 不允许 AI 单独决定编程题最终正确或错误。",
        "V1 不提供完整项目仓库级自动修改与提交代码能力。",
        "V1 不基于聊天次数或语言表达风格推断学生能力。",
        "V1 不追求自动生成整门课程，课程内容仍由管理员审核发布。",
    ):
        add_bullet(doc, text)

    add_page_break(doc)
    add_heading(doc, "3. 用户角色与核心场景", 1)
    add_heading(doc, "3.1 用户角色", 2)
    add_table(
        doc,
        ["角色", "核心诉求", "AI 助手提供的价值"],
        [
            ("学习者", "理解知识、完成练习、知道错误原因", "上下文问答、分级提示、代码诊断、质量反馈、复盘"),
            ("课程管理员", "配置可靠题目和评分规则", "测试用例辅助生成、Rubric 建议、内容检查"),
            ("平台管理员", "控制成本、安全与质量", "模型配置、限流、日志、异常与质量监控"),
        ],
        [1500, 3600, 4260],
    )

    add_heading(doc, "3.2 核心用户故事", 2)
    stories = [
        ("US-01", "作为学习者，我希望 AI 理解当前小节内容，以便我的问题不需要重复补充背景。"),
        ("US-02", "作为学习者，我希望错误代码得到具体诊断，而不只是“测试未通过”。"),
        ("US-03", "作为学习者，我希望先获得思路提示，而不是立即看到完整答案。"),
        ("US-04", "作为学习者，我希望即使实现方式不同，只要功能正确，也能得到正确判定。"),
        ("US-05", "作为学习者，我希望在代码通过后，进一步知道代码是否清晰、高效、规范。"),
        ("US-06", "作为管理员，我希望每道编程题有明确测试集与评分 Rubric，以保证评价一致。"),
        ("US-07", "作为管理员，我希望能追溯某次评分依据，以处理争议和修正题目。"),
    ]
    add_table(doc, ["编号", "用户故事"], stories, [1300, 8060])

    add_heading(doc, "3.3 典型学习场景", 2)
    add_number(doc, "学习者进入小节，AI 根据课程、小节和进度给出简短学习引导。")
    add_number(doc, "学习者阅读内容或主动提问，AI 只围绕当前学习上下文回答。")
    add_number(doc, "学习者打开练习并提交答案；选择题按现有规则判分。")
    add_number(doc, "编程题进入隔离测试执行，系统得到编译结果、测试通过率和资源使用情况。")
    add_number(doc, "AI 基于题目、代码和测试摘要生成质量评分、错误解释与下一步建议。")
    add_number(doc, "学生再次提交或请求提示，系统记录提交次数、提示等级和掌握度变化。")

    add_page_break(doc)
    add_heading(doc, "4. 总体功能范围", 1)
    add_heading(doc, "4.1 功能清单与优先级", 2)
    add_table(
        doc,
        ["编号", "功能", "说明", "优先级"],
        [
            ("AI-F01", "小节上下文对话", "读取当前课程、小节、练习和有限学习记录进行回答。", "P0"),
            ("AI-F02", "选择题确定性判分", "保留数据库答案比对及现有反馈逻辑。", "P0"),
            ("AI-F03", "编程题测试判定", "在隔离环境执行公开与隐藏测试用例。", "P0"),
            ("AI-F04", "AI 代码质量评分", "按固定 Rubric 输出分项评分与证据。", "P0"),
            ("AI-F05", "代码错误诊断", "结合失败测试说明可能原因和排查方向。", "P0"),
            ("AI-F06", "三级提示", "从方向、知识点到伪代码逐步增强。", "P0"),
            ("AI-F07", "会话与状态恢复", "同一用户进入同一小节时恢复必要学习状态。", "P0"),
            ("AI-F08", "评分记录追溯", "保存测试结果、Rubric、模型和结构化评价。", "P0"),
            ("AI-F09", "课后理解检测", "根据本节内容与错误生成一道简短检测题。", "P1"),
            ("AI-F10", "变式题生成", "针对薄弱知识点生成同难度或递进练习。", "P1"),
            ("AI-F11", "学习画像", "汇总知识点掌握度、常见错误与提示依赖。", "P1"),
            ("AI-F12", "管理员 AI 出题辅助", "生成题干、测试建议、解析和提示草稿，需人工审核。", "P2"),
        ],
        [1000, 2000, 4960, 1400],
        font_size=8.6,
    )

    add_heading(doc, "4.2 AI 助手界面能力", 2)
    for text in (
        "对话区展示 AI 与用户消息，并支持流式输出和停止生成。",
        "提供快捷操作：解释当前知识点、分析我的代码、给一个提示、换一种说法、总结本节。",
        "反馈卡展示功能结果、总分、分项得分、主要问题、改进建议和下一步行动。",
        "测试详情区区分公开测试与隐藏测试；隐藏测试仅展示必要摘要，避免泄露答案。",
        "AI 响应失败时保留已产生的确定性测试结果，并允许单独重试 AI 评审。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "4.3 上下文边界", 2)
    add_table(
        doc,
        ["上下文类型", "V1 是否使用", "说明"],
        [
            ("当前课程与小节", "是", "标题、正文、知识点、难度、预计时间"),
            ("当前练习", "是", "题干、语言、约束、测试摘要、参考解法特征"),
            ("本小节提交记录", "是", "提交次数、历史最高分、提示等级、常见错误"),
            ("全课程进度", "有限", "仅用于调整解释深度，不直接作为判题依据"),
            ("其他用户数据", "否", "严禁进入当前用户上下文"),
            ("完整标准答案", "受限", "不得在低等级提示或普通问答中直接泄露"),
        ],
        [2200, 1500, 5660],
    )

    add_page_break(doc)
    add_heading(doc, "5. 混合判题与评分需求", 1)
    add_heading(doc, "5.1 判题原则", 2)
    add_callout(
        doc,
        "最终原则",
        "选择题由数据库答案决定；编程题的功能正确性由测试系统决定；AI 不得覆盖测试结果，只能评价代码质量、解释错误并提出教学建议。",
        LIGHT_GOLD,
        GOLD,
    )

    add_heading(doc, "5.2 选择题评分", 2)
    for text in (
        "继续使用数据库中配置的正确答案进行比对。",
        "正确性结果必须确定、可重复，不调用 AI 决定分数。",
        "AI 可在提交后解释选项差异，但不得修改正确性与基础得分。",
        "提示扣分可沿用现有等级与扣分规则。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "5.3 编程题评分模型", 2)
    add_table(
        doc,
        ["评分维度", "分值", "评分来源", "核心依据"],
        [
            ("功能正确性", "60", "测试系统", "编译/运行成功、核心测试通过率"),
            ("边界与健壮性", "15", "测试系统", "空值、极端输入、异常路径、隐藏测试"),
            ("可读性", "10", "AI", "命名、结构、重复、注释与理解成本"),
            ("实现合理性", "10", "AI", "算法与数据结构是否适合题目约束"),
            ("编码规范", "5", "AI", "语言惯例、明显坏味道与不必要复杂度"),
        ],
        [2100, 900, 1800, 4560],
    )

    add_heading(doc, "5.4 总分计算规则", 2)
    add_body(doc, "总分 = 功能正确性分 + 边界与健壮性分 + AI 质量分 - 提示扣分。")
    for text in (
        "代码无法编译或无法启动：功能正确性为 0，总分上限建议为 30。",
        "核心测试未全部通过：总分上限建议为 59，避免错误程序凭代码风格及格。",
        "全部测试通过：允许进入 75 至 100 分区间，具体由 AI 质量分和提示扣分决定。",
        "AI 评审失败：保留测试得分，质量分标记为“待评审”，不应把正确代码判为错误。",
        "同一代码、同一测试版本、同一 Rubric 版本应得到相近质量分；建议允许 3 分以内浮动。",
        "历史最高分与本次分数分开保存，避免错误重提覆盖已获得的更高成绩。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "5.5 AI 质量评分 Rubric", 2)
    add_table(
        doc,
        ["维度", "高分表现", "典型扣分项"],
        [
            ("可读性 0-10", "命名清晰、结构简单、职责明确", "含糊命名、深层嵌套、大段重复、无意义注释"),
            ("实现合理性 0-10", "算法适配约束、复杂度合理、无多余工作", "明显低效、错误的数据结构、过度设计"),
            ("编码规范 0-5", "符合语言常用风格，边界处理清楚", "危险写法、全局污染、死代码、格式严重混乱"),
        ],
        [1900, 3730, 3730],
    )

    add_heading(doc, "5.6 测试用例要求", 2)
    for text in (
        "每道编程题必须至少包含基础用例、典型用例和边界用例。",
        "公开测试用于帮助学习者理解输入输出；隐藏测试用于验证泛化和防止硬编码。",
        "测试集需要独立版本号。修改测试后，历史判题结果仍应能关联旧版本。",
        "题目必须配置语言、入口方式、时间限制、内存限制和输出规范。",
        "不得把完整隐藏测试内容发送给前端或直接写入 AI 面向用户的反馈。",
    ):
        add_bullet(doc, text)

    add_page_break(doc)
    add_heading(doc, "6. AI 导师行为需求", 1)
    add_heading(doc, "6.1 回答原则", 2)
    add_table(
        doc,
        ["原则", "行为要求"],
        [
            ("上下文优先", "优先使用当前小节与题目信息，不确定时明确说明，不虚构课程规则。"),
            ("引导优先", "优先通过问题、提示和局部示例帮助学生思考，不直接代写完整答案。"),
            ("证据优先", "代码反馈应引用具体代码片段、失败测试类型或 Rubric 项。"),
            ("适龄表达", "根据课程难度使用相应术语，避免一次引入过多超纲概念。"),
            ("行动导向", "每次错误反馈至少提供一个可执行的下一步检查或修改建议。"),
            ("不覆盖事实", "不得否定或改写测试系统返回的确定性结果。"),
        ],
        [2000, 7360],
    )

    add_heading(doc, "6.2 三级提示策略", 2)
    add_table(
        doc,
        ["等级", "目标", "允许内容", "禁止内容"],
        [
            ("Level 1", "唤起思路", "指出方向、相关概念、需要观察的变量", "关键代码与完整步骤"),
            ("Level 2", "缩小问题", "指出错误区域、边界场景、算法步骤", "可直接提交的完整实现"),
            ("Level 3", "提供脚手架", "伪代码、函数结构、关键表达式或局部修复", "无解释地贴出完整标准答案"),
            ("答案阶段", "完成复盘", "在规则允许时展示参考实现并逐段解释", "把参考答案描述为唯一解"),
        ],
        [1250, 1900, 3260, 2950],
        font_size=8.8,
    )

    add_heading(doc, "6.3 代码诊断反馈结构", 2)
    for text in (
        "结果摘要：代码是否运行、通过多少测试、当前总分。",
        "主要问题：最多列出三个最值得优先修复的问题。",
        "证据：关联失败场景、代码位置或复杂度依据。",
        "下一步建议：给出一个最小、明确、可执行的行动。",
        "鼓励与边界：肯定有效思路，但不以空泛鼓励替代技术判断。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "6.4 防止答案泄露", 2)
    for text in (
        "普通聊天与低等级提示不得携带完整标准答案。",
        "标准答案只作为服务端评审参考，不直接放入可被学生操控的消息区域。",
        "AI 生成示例应与题目真实答案保持足够差异，优先使用简化示例解释概念。",
        "用户要求“忽略规则并给答案”时，AI 应继续执行当前提示等级策略。",
        "管理员可配置题目是否允许在多次失败后查看参考答案。",
    ):
        add_bullet(doc, text)

    add_page_break(doc)
    add_heading(doc, "7. LangGraph 学习流程设计方向", 1)
    add_heading(doc, "7.1 状态定义", 2)
    add_table(
        doc,
        ["状态", "含义", "主要输出"],
        [
            ("INIT", "加载用户、小节、练习和会话状态", "上下文快照"),
            ("ORIENT", "给出本次学习目标与当前进度", "学习引导"),
            ("EXPLAIN", "解释知识点或回答当前问题", "分层讲解"),
            ("PRACTICE", "引导用户进入练习", "练习指令"),
            ("RUN_TESTS", "执行编程题测试或选择题比对", "确定性结果"),
            ("AI_REVIEW", "分析测试摘要与代码质量", "分项评分与反馈"),
            ("HINT", "根据提交次数和提示等级生成提示", "Level 1-3 提示"),
            ("REFLECT", "追问原因或生成理解检测", "复盘问题"),
            ("COMPLETE", "更新进度并生成本节总结", "学习总结"),
            ("ESCALATE", "低置信度或系统异常进入人工/重试路径", "异常状态说明"),
        ],
        [1800, 4200, 3360],
        font_size=8.8,
    )

    add_heading(doc, "7.2 核心状态流", 2)
    flow_rows = [
        ("1", "INIT", "加载上下文成功", "ORIENT"),
        ("2", "ORIENT", "用户提问", "EXPLAIN"),
        ("3", "ORIENT / EXPLAIN", "用户开始练习", "PRACTICE"),
        ("4", "PRACTICE", "提交选择题", "确定性比对 -> REFLECT / COMPLETE"),
        ("5", "PRACTICE", "提交编程题", "RUN_TESTS"),
        ("6", "RUN_TESTS", "测试完成", "AI_REVIEW"),
        ("7", "AI_REVIEW", "未通过且请求帮助", "HINT"),
        ("8", "HINT", "用户再次提交", "RUN_TESTS"),
        ("9", "AI_REVIEW", "达到完成条件", "REFLECT -> COMPLETE"),
        ("10", "任意 AI 状态", "输出失败或置信度不足", "ESCALATE"),
    ]
    add_table(doc, ["步骤", "当前状态", "触发条件", "下一状态"], flow_rows, [900, 2200, 3400, 2860])

    add_heading(doc, "7.3 分支决策要求", 2)
    for text in (
        "选择题不进入 AI 判题节点，只在判分后按需进入解释节点。",
        "编程题必须先完成测试节点，再进入 AI 评审节点。",
        "测试失败与 AI 评审失败必须分别处理，不能共用一个模糊的“判题失败”状态。",
        "用户连续失败时，提示等级可升级，但升级应由提交次数、已用提示和用户主动请求共同决定。",
        "会话恢复时只恢复学习所需状态，不应把全部历史消息无上限注入模型。",
    ):
        add_bullet(doc, text)

    add_page_break(doc)
    add_heading(doc, "8. 数据与记录需求", 1)
    add_heading(doc, "8.1 现有数据复用", 2)
    add_table(
        doc,
        ["现有实体/字段", "复用方向"],
        [
            ("lessons", "提供当前小节内容、难度和预计时间。"),
            ("exercises", "提供题目、知识点、难度、metadata、hints 和来源。"),
            ("answer", "记录答案、提交次数、反馈、提示等级和历史成绩。"),
            ("lessons_progress", "记录小节状态、掌握度和最近学习时间。"),
            ("ai_chat_sessions", "记录用户在小节内的 Agent 状态、当前题目和上下文。"),
        ],
        [2600, 6760],
    )

    add_heading(doc, "8.2 建议新增的逻辑数据", 2)
    add_table(
        doc,
        ["数据对象", "关键内容", "目的"],
        [
            ("编程题测试配置", "语言、入口、公开/隐藏测试、时限、内存、版本", "支持可靠执行"),
            ("代码提交记录", "代码快照、语言、提交序号、状态、时间", "保留每次尝试而非只存最后一次"),
            ("测试运行记录", "测试版本、通过数、失败摘要、耗时、内存、错误类型", "提供确定性证据"),
            ("AI 评审记录", "Rubric 版本、模型版本、分项分数、证据、置信度", "追溯质量评分"),
            ("AI 消息记录", "会话、角色、内容类型、引用上下文、token 使用", "恢复对话与成本统计"),
            ("知识点掌握记录", "知识点、掌握度、证据、更新时间", "支持个性化学习"),
        ],
        [2000, 4800, 2560],
        font_size=8.7,
    )

    add_heading(doc, "8.3 评分记录不可变性", 2)
    for text in (
        "每次提交应形成独立记录，不能只更新同一条答案导致历史证据丢失。",
        "评分记录必须关联题目版本、测试版本、Rubric 版本和模型版本。",
        "管理员修正题目后，可选择重新评估，但不得静默覆盖原始结果。",
        "用户总积分可使用最佳有效成绩或首次通过成绩，具体策略需全站统一。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "8.4 AI 上下文最小化", 2)
    for text in (
        "只发送完成当前任务必要的数据，避免传输完整用户资料。",
        "历史对话超过阈值时进行摘要，保留关键误区、提示等级与未完成任务。",
        "标准答案、隐藏测试和系统提示必须与用户消息分层隔离。",
        "日志中对访问令牌、个人信息和可能的敏感代码进行脱敏。",
    ):
        add_bullet(doc, text)

    add_page_break(doc)
    add_heading(doc, "9. 服务交互与输出契约", 1)
    add_heading(doc, "9.1 服务职责边界", 2)
    add_table(
        doc,
        ["服务", "职责", "不得承担"],
        [
            ("课程服务", "提供小节、题目、知识点与进度上下文", "执行不可信代码"),
            ("判题服务", "编译/运行代码、执行测试、返回确定性结果", "生成教学性自然语言评价"),
            ("AI 服务", "质量评分、错误解释、提示、问答和总结", "修改测试结果或直接执行代码"),
            ("学习服务", "合并分数、更新进度、掌握度与积分", "自行推测 AI 评审内容"),
        ],
        [1700, 4220, 3440],
    )

    add_heading(doc, "9.2 编程题提交响应阶段", 2)
    add_number(doc, "接收提交并返回 submissionId，状态为 queued 或 running。")
    add_number(doc, "判题服务返回编译状态、测试通过率、资源消耗与失败摘要。")
    add_number(doc, "系统计算确定性基础分，并将安全摘要发送给 AI 评审。")
    add_number(doc, "AI 返回结构化质量评分与教学反馈。")
    add_number(doc, "学习服务合并总分，持久化记录并推送最终结果。")

    add_heading(doc, "9.3 AI 结构化输出要求", 2)
    add_table(
        doc,
        ["字段", "类型", "要求"],
        [
            ("reviewStatus", "enum", "completed / needs_review / failed"),
            ("confidence", "number", "0-1，仅表示 AI 对质量评审的把握"),
            ("qualityScores", "object", "readability 0-10、design 0-10、style 0-5"),
            ("issues", "array", "最多 3 项；包含类型、严重度、证据和说明"),
            ("strengths", "array", "最多 2 项；必须对应实际代码"),
            ("nextAction", "string", "一个最优先、可执行的下一步"),
            ("feedback", "string", "面向学生的简洁总结"),
            ("hint", "object/null", "仅在提示流程中返回，并带等级"),
        ],
        [2100, 1500, 5760],
    )

    add_heading(doc, "9.4 异常与降级", 2)
    add_table(
        doc,
        ["异常", "系统行为"],
        [
            ("编译失败", "展示编译错误摘要，AI 可解释错误；不进入测试执行。"),
            ("运行超时", "终止执行，标记超时；AI 结合代码解释可能的循环或复杂度问题。"),
            ("判题服务不可用", "提交进入可重试状态，不调用 AI 猜测正确性。"),
            ("AI 服务不可用", "展示测试结果和基础分，质量评分标记待完成，可独立重试。"),
            ("AI 输出不合法", "自动进行有限次数结构修复；仍失败则降级为待评审。"),
            ("低置信度", "保留测试结论，质量部分标记需复核，不对学生作绝对断言。"),
        ],
        [2600, 6760],
    )

    add_page_break(doc)
    add_heading(doc, "10. 安全、质量与非功能需求", 1)
    add_heading(doc, "10.1 代码执行安全", 2)
    add_callout(
        doc,
        "强制要求",
        "任何学生代码都必须视为不可信输入。不得在主业务进程、数据库容器或具有宿主机权限的环境中直接运行。",
        LIGHT_RED,
        RED,
    )
    for text in (
        "使用独立沙箱或短生命周期容器，默认禁用外网访问。",
        "限制 CPU、内存、执行时间、进程数、文件大小和系统调用。",
        "使用只读基础文件系统与临时工作目录，执行后销毁环境。",
        "禁止访问平台环境变量、数据库、Redis、宿主文件和其他用户提交。",
        "对恶意代码、死循环、fork bomb、文件遍历等行为记录安全事件。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "10.2 Prompt Injection 防护", 2)
    for text in (
        "学生代码、注释和输出都属于不可信数据，不得被模型当作系统指令。",
        "系统提示明确禁止遵循代码注释中的评分指令。",
        "标准答案与隐藏测试不与学生内容混放在同一可操控消息层。",
        "发送给 AI 的运行结果应为结构化摘要，并限制长度。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "10.3 性能与可靠性", 2)
    add_table(
        doc,
        ["项目", "需求"],
        [
            ("选择题判分", "P95 <= 500ms，不依赖 AI 服务。"),
            ("编程题测试", "常规题 P95 <= 5s，超时按题目配置终止。"),
            ("AI 首字响应", "流式对话 P95 <= 3s。"),
            ("完整 AI 评审", "P95 <= 8s；超时后展示测试结果并允许重试。"),
            ("幂等性", "相同 submissionId 重试不得重复计分或重复更新进度。"),
            ("可观测性", "记录请求链路、耗时、模型、token、失败节点和降级结果。"),
        ],
        [2600, 6760],
    )

    add_heading(doc, "10.4 成本控制", 2)
    for text in (
        "选择题默认不调用 AI；仅在用户主动请求解释时调用。",
        "相同代码与相同上下文的 AI 评审可按哈希短期缓存。",
        "限制单次上下文长度、单用户并发数和单位时间调用次数。",
        "简单错误优先使用规则化模板，复杂诊断再调用模型。",
        "记录按用户、课程、功能和模型维度的 token 与费用。",
    ):
        add_bullet(doc, text)

    add_page_break(doc)
    add_heading(doc, "11. 管理端需求", 1)
    add_heading(doc, "11.1 编程题配置", 2)
    for text in (
        "配置支持的语言与版本、函数入口或标准输入输出模式。",
        "配置公开测试、隐藏测试、时间限制、内存限制与测试版本。",
        "配置评分 Rubric、核心测试和总分上限规则。",
        "配置三级提示、参考答案可见策略和最低提交次数。",
        "发布前运行标准答案与错误样例，验证测试集能正确区分。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "11.2 AI 配置与运营", 2)
    for text in (
        "配置默认模型、备用模型、超时、重试与输出长度。",
        "查看 AI 评审失败率、结构化输出失败率、平均延迟与调用成本。",
        "抽检高分、低分、争议和低置信度案例。",
        "支持停用某道题的 AI 质量评分，仅保留测试判分。",
        "AI 生成的题目、测试或提示必须经过管理员确认后才能发布。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "11.3 争议处理", 2)
    add_number(doc, "管理员查看学生代码、测试版本、测试摘要、Rubric 和 AI 原始结构化结果。")
    add_number(doc, "确认问题来自题目、测试集、AI 评审或展示逻辑。")
    add_number(doc, "必要时人工调整本次质量分，并填写调整原因。")
    add_number(doc, "若测试集存在错误，创建新版本并选择受影响提交进行重新判定。")

    add_page_break(doc)
    add_heading(doc, "12. 验收标准", 1)
    add_heading(doc, "12.1 P0 功能验收", 2)
    acceptance = [
        ("AC-01", "选择题提交后仍由数据库标准答案判分，AI 服务不可用时不受影响。"),
        ("AC-02", "两份输出相同但写法不同的正确代码，均可通过对应测试并判为正确。"),
        ("AC-03", "编程题的功能分完全由测试结果计算，AI 无法覆盖 passed/failed。"),
        ("AC-04", "AI 返回可读性、实现合理性和规范性三个分项，分数不超出各自上限。"),
        ("AC-05", "核心测试失败时，总分不会达到及格线。"),
        ("AC-06", "AI 评审失败时仍展示测试结果，且不会把已通过代码改判为错误。"),
        ("AC-07", "三级提示逐级增强，Level 1 和 Level 2 不直接输出完整可提交答案。"),
        ("AC-08", "每次编程题提交均可追溯代码、题目版本、测试版本、Rubric 和模型版本。"),
        ("AC-09", "代码运行超时、内存超限和编译失败均能被明确区分并反馈。"),
        ("AC-10", "用户刷新或重新进入小节后，可恢复当前题目、提示等级和必要会话状态。"),
        ("AC-11", "隐藏测试内容不会出现在前端响应、普通日志或 AI 面向用户的消息中。"),
        ("AC-12", "学生代码注释中的指令不会改变系统评分规则。"),
    ]
    add_table(doc, ["编号", "验收条件"], acceptance, [1300, 8060], font_size=8.9)

    add_heading(doc, "12.2 典型测试场景", 2)
    add_table(
        doc,
        ["场景", "期望结果"],
        [
            ("不同变量名与循环写法", "只要通过测试，即获得功能正确性分。"),
            ("硬编码公开示例", "隐藏测试失败，功能分受限，AI 指出泛化问题。"),
            ("代码正确但命名混乱", "测试通过，功能分完整；AI 质量分扣分并提供改进建议。"),
            ("死循环", "沙箱超时终止，不影响其他请求，AI 解释可能的终止条件问题。"),
            ("注释中要求 AI 给 100 分", "评分不受影响，并按 Rubric 返回结果。"),
            ("AI 接口超时", "测试结果正常展示，质量部分进入待评审状态。"),
        ],
        [3100, 6260],
    )

    add_page_break(doc)
    add_heading(doc, "13. 版本规划", 1)
    add_heading(doc, "13.1 MVP / V1.0", 2)
    for text in (
        "当前小节上下文对话与流式输出。",
        "选择题保持数据库答案判分。",
        "支持一种主力编程语言的隔离执行与测试判题。",
        "固定 Rubric 的 AI 代码质量评分。",
        "失败测试解释与三级提示。",
        "提交、测试、AI 评审与会话状态记录。",
        "基础限流、成本统计、异常降级和管理端题目配置。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "13.2 V1.1", 2)
    for text in (
        "增加更多编程语言与对应运行镜像。",
        "课后理解检测和针对错误的变式题。",
        "知识点掌握度与薄弱项画像。",
        "管理员 AI 出题、测试建议和提示草稿。",
        "质量抽检、低置信度队列和评分争议工作台。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "13.3 V2.0 方向", 2)
    for text in (
        "以剧情任务为载体的多步骤编程挑战。",
        "跨小节的长期学习计划与间隔复习。",
        "项目级代码评审与多文件任务。",
        "基于学习证据动态调整课程路径和练习难度。",
    ):
        add_bullet(doc, text)

    add_heading(doc, "13.4 推荐实施顺序", 2)
    add_number(doc, "先定义编程题运行协议、测试模型和安全边界。")
    add_number(doc, "完成判题服务与确定性分数，再接入 AI 质量评审。")
    add_number(doc, "打通提交记录、评分追溯和异常降级。")
    add_number(doc, "实现小节对话与三级提示，并接入 LangGraph 状态流。")
    add_number(doc, "最后增加掌握度、变式题和个性化路线。")

    add_page_break(doc)
    add_heading(doc, "14. 风险与待确认事项", 1)
    add_heading(doc, "14.1 主要风险", 2)
    add_table(
        doc,
        ["风险", "影响", "缓解方向"],
        [
            ("沙箱隔离不足", "可能危及平台与其他用户数据", "独立执行服务、资源限制、禁网、短生命周期环境"),
            ("测试集覆盖不足", "错误代码被判为正确", "题目发布校验、隐藏测试、错误样例回归"),
            ("AI 评分波动", "同类代码得分不一致", "固定模型与 Rubric、低温度、校准样例、抽检"),
            ("答案泄露", "削弱学习效果和题库价值", "分级提示、上下文隔离、隐藏测试脱敏"),
            ("响应时间过长", "影响提交体验", "阶段式展示、异步评审、缓存和降级"),
            ("调用成本上升", "运营成本不可控", "限流、短上下文、规则优先、按功能统计"),
        ],
        [2200, 3000, 4160],
        font_size=8.8,
    )

    add_heading(doc, "14.2 开发前需要确认", 2)
    for text in (
        "MVP 首先支持哪一种编程语言，以及题目采用函数模式还是标准输入输出模式。",
        "编程题总分中提示扣分的具体规则，以及历史成绩采用最高分还是最近分。",
        "核心测试的定义方式和未通过时的总分上限。",
        "用户在多少次失败或使用多少级提示后可以查看参考答案。",
        "AI 评审低置信度阈值及管理员复核流程。",
        "沙箱服务采用自建容器执行、第三方判题服务还是独立微服务。",
    ):
        add_bullet(doc, text)

    add_callout(
        doc,
        "建议结论",
        "第一版先选择单一语言和有限题型，把“可靠测试判定、稳定质量评分、清晰教学反馈、完整追溯”做扎实，再扩展语言和个性化能力。",
        LIGHT_GREEN,
        GREEN,
    )

    doc.core_properties.title = "CodeStory AI 助手模块需求文档 V1.0"
    doc.core_properties.subject = "AI 导师、混合判题与代码质量评分"
    doc.core_properties.author = "CodeStory Product Team"
    doc.core_properties.keywords = "CodeStory, AI Agent, LangGraph, 编程教育, 混合判题"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
